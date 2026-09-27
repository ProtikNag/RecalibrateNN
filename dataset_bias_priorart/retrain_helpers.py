import copy
import os
import time
from dataclasses import asdict
from typing import Any, Dict, List, Optional, Tuple

import numpy as np
import pandas as pd
import torch
import torch.nn as nn
from torch.utils.data import DataLoader


def extract_logits(model_output: Any) -> torch.Tensor:
    if isinstance(model_output, torch.Tensor):
        return model_output
    if hasattr(model_output, "logits"):
        return model_output.logits
    if isinstance(model_output, (list, tuple)) and len(model_output) > 0:
        return model_output[0]
    raise TypeError("Unsupported model output format while extracting logits")


def infer_sample_paths_and_labels(dataset: Any) -> Tuple[List[str], List[int]]:
    if hasattr(dataset, "samples"):
        samples = getattr(dataset, "samples")
        paths = [str(x[0]) for x in samples]
        labels = [int(x[1]) for x in samples]
        return paths, labels

    if hasattr(dataset, "getfilelist"):
        paths, labels = dataset.getfilelist()
        return [str(p) for p in paths], [int(y) for y in labels]

    raise ValueError("Unable to infer sample paths from dataset. Expected .samples or .getfilelist().")


def infer_group_labels_from_parent_folder(paths: List[str]) -> Tuple[np.ndarray, Dict[int, str]]:
    parent_names = [os.path.basename(os.path.dirname(p)) for p in paths]
    uniq = sorted(set(parent_names))
    group_to_id = {name: idx for idx, name in enumerate(uniq)}
    labels = np.array([group_to_id[name] for name in parent_names], dtype=np.int64)
    id_to_group = {v: k for k, v in group_to_id.items()}
    return labels, id_to_group


def make_loader(
    dataset: Any,
    batch_size: int,
    shuffle: bool,
    seed: int,
    num_workers: int,
    pin_memory: bool,
) -> DataLoader:
    generator = torch.Generator()
    generator.manual_seed(seed)
    return DataLoader(
        dataset,
        batch_size=batch_size,
        shuffle=shuffle,
        num_workers=num_workers,
        generator=generator,
        pin_memory=pin_memory,
    )


def train_one_method(
    model: nn.Module,
    train_loader: DataLoader,
    valid_loader: DataLoader,
    device: str,
    epochs: int,
    learning_rate: float,
    max_grad_norm: float = 7.0,
) -> Dict[str, float]:
    model = model.to(device)
    model.train()
    optimizer = torch.optim.Adam(model.parameters(), lr=learning_rate)
    loss_fn = nn.CrossEntropyLoss()

    final_train_loss = 0.0

    for _epoch in range(epochs):
        running = 0.0
        batches = 0
        for imgs, labels in train_loader:
            imgs = imgs.to(device)
            labels = labels.to(device)
            optimizer.zero_grad()
            logits = extract_logits(model(imgs))
            loss = loss_fn(logits, labels)
            loss.backward()
            torch.nn.utils.clip_grad_norm_(model.parameters(), max_norm=max_grad_norm)
            optimizer.step()
            running += float(loss.item())
            batches += 1
        final_train_loss = running / max(1, batches)

    valid_metrics = evaluate_model(model, valid_loader, device)
    valid_metrics["train_loss"] = final_train_loss
    return valid_metrics


def evaluate_model(model: nn.Module, loader: DataLoader, device: str) -> Dict[str, float]:
    model.eval()
    loss_fn = nn.CrossEntropyLoss()
    correct = 0
    total = 0
    running = 0.0
    batches = 0

    with torch.no_grad():
        for imgs, labels in loader:
            imgs = imgs.to(device)
            labels = labels.to(device)
            logits = extract_logits(model(imgs))
            loss = loss_fn(logits, labels)
            preds = torch.argmax(logits, dim=1)
            correct += int((preds == labels).sum().item())
            total += int(labels.size(0))
            running += float(loss.item())
            batches += 1

    return {
        "valid_loss": running / max(1, batches),
        "valid_accuracy": float(correct) / float(total) if total > 0 else 0.0,
    }


def sanitize_sheet_name(name: str, used_names: set) -> str:
    invalid_chars = set('[]:*?/\\')
    cleaned = "".join("_" if c in invalid_chars else c for c in str(name))
    cleaned = cleaned[:31] if cleaned else "Sheet"
    base = cleaned
    suffix = 1
    while cleaned.lower() in used_names:
        tail = str(suffix)
        cleaned = (base[: 31 - len(tail)] + tail)[:31]
        suffix += 1
    used_names.add(cleaned.lower())
    return cleaned


def style_worksheet(ws: Any) -> None:
    try:
        from openpyxl.styles import Alignment, Font, PatternFill
    except Exception:
        return

    if ws.max_row < 1:
        return

    fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    font = Font(bold=True, color="FFFFFF")

    for cell in ws[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            value = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(value))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 50)


def result_to_detail_rows(name: str, result: Any) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    if result.group_losses is not None:
        for gid, loss in sorted(result.group_losses.items()):
            rows.append({"method": name, "detail_type": "group_loss", "group_id": int(gid), "value": float(loss)})

    if result.bias_scores is not None:
        for idx, score in enumerate(result.bias_scores.tolist()):
            rows.append({"method": name, "detail_type": "bias_score", "index": idx, "value": float(score)})

    if result.subgroups is not None:
        for idx, subgroup in enumerate(result.subgroups):
            rows.append(
                {
                    "method": name,
                    "detail_type": "subgroup",
                    "subgroup_id": idx,
                    "subgroup_size": int(len(subgroup)),
                    "indices": ",".join(str(int(x)) for x in subgroup[:300]),
                }
            )

    if getattr(result, "details", None):
        for d_key, d_val in result.details.items():
            rows.append({"method": name, "detail_type": "meta", "key": d_key, "value": str(d_val)})

    return rows


def write_retraining_excel(
    output_path: str,
    run_params: Dict[str, Any],
    group_mapping: Dict[int, str],
    detection_results: Dict[str, Any],
    training_rows: List[Dict[str, Any]],
    failed_rows: List[Dict[str, Any]],
) -> None:
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    used_names: set = set()

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        run_df = pd.DataFrame([{"key": k, "value": str(v)} for k, v in run_params.items()])
        run_df.to_excel(writer, sheet_name=sanitize_sheet_name("Run_Parameters", used_names), index=False)

        map_df = pd.DataFrame(
            [{"group_id": gid, "group_name": gname} for gid, gname in sorted(group_mapping.items())]
        )
        map_df.to_excel(writer, sheet_name=sanitize_sheet_name("Group_Mapping", used_names), index=False)

        detect_summary = []
        for key, result in detection_results.items():
            scores = result.bias_scores if result.bias_scores is not None else np.array([], dtype=np.float64)
            detect_summary.append(
                {
                    "method": key,
                    "num_scores": int(len(scores)),
                    "score_mean": float(np.mean(scores)) if len(scores) else np.nan,
                    "score_max": float(np.max(scores)) if len(scores) else np.nan,
                    "score_min": float(np.min(scores)) if len(scores) else np.nan,
                    "num_group_losses": int(len(result.group_losses or {})),
                    "num_subgroups": int(len(result.subgroups or [])),
                }
            )
        pd.DataFrame(detect_summary).to_excel(
            writer,
            sheet_name=sanitize_sheet_name("Method_Detection_Summary", used_names),
            index=False,
        )

        pd.DataFrame(training_rows).to_excel(
            writer,
            sheet_name=sanitize_sheet_name("Method_Training_Summary", used_names),
            index=False,
        )

        if failed_rows:
            pd.DataFrame(failed_rows).to_excel(
                writer,
                sheet_name=sanitize_sheet_name("Failed_Methods", used_names),
                index=False,
            )

        for method_name, result in detection_results.items():
            detail_rows = result_to_detail_rows(method_name, result)
            if detail_rows:
                pd.DataFrame(detail_rows).to_excel(
                    writer,
                    sheet_name=sanitize_sheet_name(f"{method_name}_details", used_names),
                    index=False,
                )

        for ws in writer.book.worksheets:
            style_worksheet(ws)

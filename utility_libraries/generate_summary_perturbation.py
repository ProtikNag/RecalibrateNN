import os
import sys
import pandas as pd
from pathlib import Path
import openpyxl
from openpyxl import Workbook

"""

"""
#3 Parse all excel files and create a list
def get_allexcelfiles(directory):
    excel_files = [f for f in os.listdir(directory) if f.endswith(('.xlsx', '.xls'))]
    gaussian_files = [f for f in excel_files if 'gaussian' in f.lower()]
    mean_files = [f for f in excel_files if 'mean' in f.lower()]
    alpha_files = [f for f in excel_files if 'alpha' in f.lower()]
    return gaussian_files, mean_files, alpha_files
def get_header():
    headers = ['File','Method Used','Alpha value', 'Artifact used', 'Class Used', 'Combination', 'Combination Layer', 'Class_ID',  'Mean_Original_Prob','Mean_Perturbed_Prob',
              'mean_original_logits_class0', 'mean_original_logits_class1', 'mean_original_logits_class2',
              'mean_perturbed_logits_class0', 'mean_perturbed_logits_class1', 'mean_perturbed_logits_class2',
               'Mean_Delta_Logits_Class0', 'Mean_Delta_Logits_Class1', 'Mean_Delta_Logits_Class2']
    return headers

# 4. Create destination file
def write_summary_to_excel(summary_data, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    # Write headers
    headers = get_header()
    ws.append(headers)
    
    # Write summary data
    for row in summary_data:
        ws.append(row)
    
    # Save the workbook
    wb.save(output_path)
    print(f"Summary saved to: {output_path}")

def create_workbook_sheets(filename):
    wb = Workbook()
    if 'Sheet' in wb.sheetnames:
        del wb['Sheet']
    for class_id in range(3):
        for artifact in ['bg', 'coat', 'face', 'legs']:
            sheet_name = f"Summary_class{class_id}_{artifact}"
            wb.create_sheet(sheet_name)
    headers = get_header()
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        ws.append(headers)
    wb.save(filename)
    wb.close()
    return

def create_desitnation_file(directory):
    temp = directory.split('/')[-1]
    parent_directory = os.path.dirname(directory)
    gaussian_summary_positive_path = os.path.join(parent_directory, temp + '_gaussian_summary.xlsx')
    mean_summary_positive_path = os.path.join(parent_directory, temp + '_mean_summary.xlsx')
    alpha_summary_positive_path = os.path.join(parent_directory, temp + '_alpha_summary.xlsx')

    create_workbook_sheets(gaussian_summary_positive_path)
    create_workbook_sheets(mean_summary_positive_path)
    create_workbook_sheets(alpha_summary_positive_path)

    gaussian_summary_negative_path = os.path.join(parent_directory, temp + '_gaussian_summary_negative.xlsx')
    mean_summary_negative_path = os.path.join(parent_directory, temp + '_mean_summary_negative.xlsx')
    alpha_summary_negative_path = os.path.join(parent_directory, temp + '_alpha_summary_negative.xlsx')
    
    create_workbook_sheets(gaussian_summary_negative_path)
    create_workbook_sheets(mean_summary_negative_path)
    create_workbook_sheets(alpha_summary_negative_path)

    result_paths = {
        'gaussian': {
            'positive': gaussian_summary_positive_path,
            'negative': gaussian_summary_negative_path
        },
        'mean': {
            'positive': mean_summary_positive_path,
            'negative': mean_summary_negative_path
        },
        'alpha': {
            'positive': alpha_summary_positive_path,
            'negative': alpha_summary_negative_path
        }
    }
    return result_paths    

def reorder_sheets(wb):
    # Get all sheet names
    all_sheets = wb.sheetnames
    
    # Define the order of artifacts
    artifact_order = ['_bg', '_coat', '_legs', '_face']
    
    # Create a list to hold sheets in the desired order
    ordered_sheets = []
    
    # Reorder sheets based on artifact suffix
    for suffix in artifact_order:
        matching_sheets = [sheet for sheet in all_sheets if sheet.endswith(suffix)]
        ordered_sheets.extend(sorted(matching_sheets))
    
    # Reorder the sheets in the workbook
    for index, sheet_name in enumerate(ordered_sheets):
        wb.move_sheet(sheet_name, offset=index - wb.sheetnames.index(sheet_name))
    return
                    
                    
def main():
    # 1. Get directory name from command line
    if len(sys.argv) < 2:
        print("Usage: python generate_summary_perturbation.py <directory_path>")
        sys.exit(1)
    
    directory = sys.argv[1]
    
    if not os.path.exists(directory):
        print(f"Error: Directory '{directory}' does not exist")
        sys.exit(1)
    # 2. Parse all excel files and create a list
    gaussian_files, mean_files, alpha_files = get_allexcelfiles(directory)
    print(f"Gaussian files: {gaussian_files}")
    print(f"Mean files: {mean_files}")
    print(f"Alpha files: {alpha_files}")
    
    # 4. Create destination file
    summary_files = create_desitnation_file(directory)
    def process_subfiles(filelist,method,  directory, destination_file):
        wb = openpyxl.load_workbook(destination_file)
        process_files(filelist, method, directory, wb, destination_file)
        reorder_sheets(wb)
        wb.save(destination_file)
        wb.close()

    positive_summary_file = summary_files['gaussian']['positive']  # Example for accessing positive summary file
    negative_summary_file = summary_files['gaussian']['negative']  # Example for accessing negative summary file
    gaussian_negative_files = [f for f in gaussian_files if 'negative' in f.lower()]
    gaussian_positive_files = [f for f in gaussian_files if 'negative' not in f.lower()]

    process_subfiles(gaussian_positive_files, 'gaussian', directory, positive_summary_file)
    process_subfiles(gaussian_negative_files, 'gaussian', directory, negative_summary_file)

    positive_summary_file = summary_files['mean']['positive']  # Example for accessing positive summary file
    negative_summary_file = summary_files['mean']['negative']  # Example for accessing negative summary file

    mean_negative_files = [f for f in mean_files if 'negative' in f.lower()]
    mean_positive_files = [f for f in mean_files if 'negative' not in f.lower()]
    process_subfiles(mean_positive_files, 'mean', directory, positive_summary_file)
    process_subfiles(mean_negative_files, 'mean', directory, negative_summary_file)

    positive_summary_file = summary_files['alpha']['positive']  # Example for accessing positive summary file
    negative_summary_file = summary_files['alpha']['negative']  # Example for accessing negative summary file
    
    alpha_negative_files = [f for f in alpha_files if 'negative' in f.lower() and '0.67' in f.lower()]
    alpha_positive_files = [f for f in alpha_files if 'negative' not in f.lower() and '0.67' in f.lower()]
    process_subfiles(alpha_positive_files, 'alpha', directory, positive_summary_file)
    process_subfiles(alpha_negative_files, 'alpha', directory, negative_summary_file)

    
    exit()

    process_subfiles(gaussian_positive_files, 'gaussian', directory, positive_summary_file)
    process_subfiles(mean_positive_files, 'mean', directory, positive_summary_file)
    process_subfiles(alpha_positive_files, 'alpha', directory, positive_summary_file)
    process_subfiles(gaussian_negative_files, 'gaussian', directory, negative_summary_file)
    process_subfiles(mean_negative_files, 'mean', directory, negative_summary_file)
    process_subfiles(alpha_negative_files, 'alpha', directory, negative_summary_file)
    return
    
    
    

    
    


    
    
def process_files(perturbation_list,method_name, directory, wb,destination_file):
    for filename in  perturbation_list:
        file_path = os.path.join(directory, filename)
        # List all sheets in the excel file
        xl_file = pd.ExcelFile(file_path)
        sheet_names = xl_file.sheet_names
        sheet_names.pop(0)  # Remove Summary sheet
        print(f"Sheets in {filename}: {sheet_names}") 
        if(method_name == "alpha"):
            alpha_value = filename.split('_')[-1].replace('.xlsx', '')
        
        method_used = method_name
        alpha_value = alpha_value if method_name == "alpha" else 'N/A'
        #Gaussian Mean and Alpha are only 3 methods considered 
        print(f"Processing file: {filename} with method: {method_used} and alpha value: {alpha_value}")
        artifact_used = filename.split('_')[-2] if method_name in ['gaussian', 'mean'] else filename.split('_')[-3]
        class_used = filename.split('_')[-3]if method_name in ['gaussian', 'mean'] else filename.split('_')[-4]
        if(artifact_used not in ['bg', 'coat', 'face', 'legs']):
            print(f"Warning: Unexpected artifact '{artifact_used}' in filename '{filename}'")
            artifact_used = class_used
            class_used = filename.split('_')[-4] if method_name in ['gaussian', 'mean'] else filename.split('_')[-5]
        
        
        # Read the summary sheet and create a list of the first column
        try:
            summary_df = pd.read_excel(file_path, sheet_name='Summary')
            combination_layernames = summary_df.iloc[:, 0].tolist()
            print(f"First column values from Summary sheet: {combination_layernames}")
            print("Number of sheets in the file:", len(sheet_names))
        except Exception as e:
            print(f"Could not read Summary sheet from {filename}: {str(e)}")
        for worksheet, combination_layer in zip(sheet_names, combination_layernames):
            try:
                # 6. Navigate to sheet name 'combination'
                df = pd.read_excel(file_path, sheet_name=worksheet)
                # 7. Data already read into dataframe
                # 8. Group by class_id (0, 1, 2)
                # Replace class_id values: 'd' -> 0, 'h' -> 1, 'z' -> 2
                #mapping = {'d': 0, 'h': 1, 'z': 2}
                #df['class_id'] = df['class_id'].map(mapping).fillna(df['class_id'])
                
                for class_id in [0, 1, 2]:
                    class_df = df[df['class_id'] == class_id].copy()
                    if class_df.empty:
                        continue
                    
                    # 9. Compute mean of original_prob
                    mean_original_prob = class_df['original_prob'].mean()
                   
                    # 10. Compute mean of perturbed_prob
                    mean_perturbed_prob = class_df['perturbed_prob'].mean()
                    #rename the column name 'oritinal_logits_tensor_class0' to 'original_logits_tensor_class0'
                    try:
                        class_df.rename(columns={'oritinal_logits_tensor_class0': 'original_logits_tensor_class0'}, inplace=True)
                    except Exception as e:
                        print(f"Error renaming column in {filename}: {str(e)}")
                    try:
                        class_df.rename(columns={'oritinal_logits_tensor_class1': 'original_logits_tensor_class1'}, inplace=True)
                    except Exception as e:
                        print(f"Error renaming column in {filename}: {str(e)}")
                    try:
                        class_df.rename(columns={'oritinal_logits_tensor_class2': 'original_logits_tensor_class2'}, inplace=True)
                    except Exception as e:
                        print(f"Error renaming column in {filename}: {str(e)}")
                        continue
                    # 11. Compute mean of delta_logits columns
                    original_logits_class0 = class_df['original_logits_tensor_class0'].mean()
                    original_logits_class1 = class_df['original_logits_tensor_class1'].mean()
                    original_logits_class2 = class_df['original_logits_tensor_class2'].mean()
                    perturbed_logits_tensor_class0 = class_df['perturbed_logits_tensor_class0'].mean()
                    perturbed_logits_tensor_class1 = class_df['perturbed_logits_tensor_class1'].mean()
                    perturbed_logits_tensor_class2 = class_df['perturbed_logits_tensor_class2'].mean()
                    
                    mean_delta_class0 = class_df['delta_logits_class0'].mean()
                    mean_delta_class1 = class_df['delta_logits_class1'].mean()
                    mean_delta_class2 = class_df['delta_logits_class2'].mean()
                    
                    # 12 & 13. Store results in summary file
                    row_data = [
                        filename,
                        method_used,
                        alpha_value,
                        artifact_used,
                        class_used,
                        worksheet,
                        combination_layer,
                        class_id,
                        mean_original_prob,
                        mean_perturbed_prob,
                        original_logits_class0,
                        original_logits_class1,
                        original_logits_class2,
                        perturbed_logits_tensor_class0,
                        perturbed_logits_tensor_class1,
                        perturbed_logits_tensor_class2,
                        mean_delta_class0,
                        mean_delta_class1,
                        mean_delta_class2
                    ]
                    ws = wb[f"Summary_class{class_id}_{artifact_used}"]
                    ws.append(row_data)
                    
                    print(f"Processed {filename} - Sheet {worksheet} - Method {method_name} - Class {class_id} - Artifact {artifact_used}")
            except Exception as e:
                print(f"Error processing {filename}: {str(e)}")
                continue

if __name__ == "__main__":
    main()
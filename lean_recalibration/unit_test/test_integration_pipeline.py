"""
End-to-end pipeline tests chaining ALL FOUR target scripts together using REAL output from
each stage feeding into the next (no synthetic/hand-crafted intermediate files) -- this is the
"whole workflow actually works" complement to the per-script test files, which each already
test one script thoroughly in isolation:
    test_main_store_cav.py
    test_main_recalib_custom_by_loading_cav.py
    test_utils_sensitivity_multiclass.py
    test_bottleneck_detection.py

Natural pipeline order
-----------------------
    main_store_cav.py
        -> main_recalib_custom_by_loading_cav.py   (consumes the CAV store)
        -> utils_sensitivity_multiclass.py          (consumes the same CAV store's manifest)
               -> bottleneck_detection.py            (consumes the sensitivity Excel report)

Multiclass mode chains all FOUR stages for real (TestFullPipelineMulticlass below), because
main_store_cav.py's --store_multiconcept_cav path and utils_sensitivity_multiclass.py's
build_concept_specs() already agree on the same class-prefixed concept-naming convention.

Single-class mode is exercised only through stage 2 (TestPartialPipelineSingleClass below).
Continuing single-class mode into stage 3 is deliberately NOT done here: doing so against a
REAL main_store_cav.py manifest was confirmed (this session) to always collect zero rows, due
to a separate, pre-existing concept-naming mismatch between the two scripts in single mode
(main_store_cav.py keys single-mode CAVs by the bare concept-folder name, e.g. "stripes", while
utils_sensitivity_multiclass.py's build_concept_specs() always looks them up under a
class-prefixed name, e.g. "cat_stripes"). Per explicit maintainer direction this test suite
documents that behaviour (see
test_utils_sensitivity_multiclass.py::TestUtilsSensitivityCli::
test_single_class_end_to_end_documents_concept_naming_mismatch) rather than modifying
production code to fix it, so it is not re-chained here into a 4-stage single-class pipeline
that would trivially produce nothing to hand to bottleneck_detection.py.
"""

import os

import openpyxl
import pytest

from helpers import constants as C


@pytest.mark.integration
@pytest.mark.slow
class TestFullPipelineMulticlass:
    def test_store_recalib_sensitivity_bottleneck_chain(self, run_script, tmp_path, populated_multi_cav_store):
        # ---- Stage 1: CAV store (already built once per session by the
        # populated_multi_cav_store fixture -- main_store_cav.py --store_multiconcept_cav) ----
        manifest_path = os.path.join(populated_multi_cav_store, f"{C.MODEL_NAME}_manifest.json")
        assert os.path.isfile(manifest_path), "Stage 1 (main_store_cav.py) did not produce a per-model manifest"

        # ---- Stage 2: joint multiclass recalibration, consuming stage 1's CAV store ----
        recalib_results = str(tmp_path / "recalib_results")
        recalib_result = run_script(
            "main_recalib_custom_by_loading_cav.py",
            [
                "--model_name", C.MODEL_NAME,
                "--model_path", C.MULTI_MODEL_ROOT,
                "--config_file", C.MULTI_CONFIG_RECALIB_PATH,
                "--cav_store", populated_multi_cav_store,
                "--store_results", recalib_results,
                "--multiclass_recalibration_mode",
            ],
            timeout=600,
        )
        recalib_result.assert_success()
        recalib_model_dir = os.path.join(recalib_results, C.MODEL_NAME)
        recalib_pth_files = [
            f for f in os.listdir(recalib_model_dir)
            if f.startswith("model_multiclass_combo") and f.endswith(".pth")
        ]
        assert recalib_pth_files, f"Stage 2 (recalibration) produced no checkpoints in {recalib_model_dir}"

        # ---- Stage 3: sensitivity computation, reading CAVs from stage 1's manifest ----
        sensitivity_results = str(tmp_path / "sensitivity_results")
        sensitivity_result = run_script(
            "utils_sensitivity_multiclass.py",
            [
                "--config", C.MULTI_CONFIG_PATH,
                "--org_model_path", os.path.join(C.MULTI_MODEL_ROOT, C.MODEL_NAME, f"{C.MODEL_NAME}.pth"),
                "--model_name", C.MODEL_NAME,
                "--store_results", sensitivity_results,
                "--manifest", manifest_path,
                "--concept_mode", "multiclass",
            ],
            timeout=600,
        )
        sensitivity_result.assert_success()
        assert "Concept mode resolved to: multiclass" in sensitivity_result.stdout
        sensitivity_model_dir = os.path.join(sensitivity_results, C.MODEL_NAME)
        xlsx_files = [
            f for f in os.listdir(sensitivity_model_dir)
            if f.startswith("sensitivity_report_") and f.endswith(".xlsx")
        ]
        assert xlsx_files, f"Stage 3 (sensitivity) produced no Excel report in {sensitivity_model_dir}"
        sensitivity_xlsx = os.path.join(sensitivity_model_dir, xlsx_files[0])

        # ---- Stage 4: bottleneck detection, consuming stage 3's REAL Excel report directly ----
        bottleneck_output = str(tmp_path / "bottleneck_report.xlsx")
        bottleneck_result = run_script(
            "bottleneck_detection.py",
            [sensitivity_xlsx, bottleneck_output, "--no-charts"],
            timeout=300,
        )
        bottleneck_result.assert_success()
        assert os.path.isfile(bottleneck_output), "Stage 4 (bottleneck_detection) produced no output workbook"

        wb = openpyxl.load_workbook(bottleneck_output, read_only=True)
        assert "Overview" in wb.sheetnames
        expected_concepts = set()
        for class_idx, folders in C.MULTI_CLASS_CONCEPTS.items():
            class_name = C.MULTI_CLASSES[class_idx]
            for folder in folders:
                expected_concepts.add(f"{class_name}_{folder}")
        for concept in expected_concepts:
            assert any(sn.startswith(concept[:20]) for sn in wb.sheetnames), (
                f"No bottleneck sheet found for concept '{concept}' among {wb.sheetnames}"
            )


@pytest.mark.integration
@pytest.mark.slow
class TestPartialPipelineSingleClass:
    def test_store_then_recalib_chain(self, run_script, tmp_path, populated_single_cav_store):
        """Single-class pipeline, stages 1-2 only (main_store_cav.py -> independent per-class
        recalibration). See the module docstring above for why this does not continue into
        stage 3/4 -- that exact single-class limitation is covered by a dedicated test in
        test_utils_sensitivity_multiclass.py instead."""
        manifest_path = os.path.join(populated_single_cav_store, f"{C.MODEL_NAME}_manifest.json")
        assert os.path.isfile(manifest_path), "Stage 1 (main_store_cav.py) did not produce a per-model manifest"

        recalib_results = str(tmp_path / "recalib_results")
        recalib_result = run_script(
            "main_recalib_custom_by_loading_cav.py",
            [
                "--model_name", C.MODEL_NAME,
                "--model_path", C.SINGLE_MODEL_ROOT,
                "--config_file", C.SINGLE_CONFIG_RECALIB_PATH,
                "--cav_store", populated_single_cav_store,
                "--store_results", recalib_results,
            ],
            timeout=600,
        )
        recalib_result.assert_success()
        recalib_model_dir = os.path.join(recalib_results, C.MODEL_NAME)
        assert os.path.isfile(os.path.join(recalib_model_dir, f"recalibration_summary_{C.MODEL_NAME}.xlsx"))
        pth_files = [f for f in os.listdir(recalib_model_dir) if f.startswith("model_cls") and f.endswith(".pth")]
        assert pth_files, f"Stage 2 (recalibration) produced no per-class checkpoints in {recalib_model_dir}"

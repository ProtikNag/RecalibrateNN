"""
bottleneck_detection.py
========================
Concept Sensitivity Propagation Index (CSPI) + layer-wise bottleneck detection for
TCAV/sensitivity reports produced by ``utils_sensitivity_multiclass.py``.

INPUT
-----
The primary input is the multiclass sensitivity **Excel report** (e.g.
``sensitivity_report_resnet50_<timestamp>.xlsx``). That workbook has one sheet per
concept (e.g. ``deer_coat``, ``deer_legs``, ``horse_background`` ...) in LONG format
(one row per sample/layer/stage), plus ``Run_Info`` / ``Summary`` / ``Summary_By_Layer``
meta sheets. This script:

  1. Reads the workbook, discovers every concept sheet (via the ``Summary`` sheet's
     ``excel_sheet_names`` column, falling back to auto-detection by header signature).
  2. Pivots each concept's long rows into a wide sample x layer sensitivity-score matrix,
     one per ``stage`` present (``before``, and ``after`` for every ``lambda_align`` found -
     this is the "flexibility to incorporate a before/after comparison based on the
     parameter present in [the] stage [column]").
  3. Computes, per concept and per stage:
        - CSPI_M : magnitude Concept-Sensitivity-Propagation-Index (mean |alignment| of a
                   layer with every downstream layer - how strongly the concept signal
                   keeps showing up later, regardless of sign).
        - CSPI_D : directional CSPI (mean *signed* alignment with every downstream layer -
                   positive means the concept keeps pointing the same way; negative flags
                   a sign flip).
        - Bottleneck flags from two complementary heuristics (global downstream alignment,
          and adjacent-layer-pair correlation drops), combined into one verdict.
  4. Builds one results **table per concept** (e.g. "a table for deer coat", "a table for
     deer legs", ...), a correlation-heatmap sheet, and (when both a 'before' and an 'after'
     stage exist for a concept) a before/after comparison table + charts.
  5. Writes everything into a single output ``.xlsx`` (tables + embedded PNG charts), plus
     the same PNG charts saved under ``<output_dir>/charts/<concept>/`` and a cross-concept
     ``Overview`` + ``Bottleneck_Summary`` (+ ``Comparison_Summary``) sheets.

A legacy wide-format CSV (one row per sample, one ``sensitivityscore_before_<layer>`` column
per layer - the format this script originally accepted) is still supported for backward
compatibility; just pass a ``.csv`` file instead of ``.xlsx``.

USAGE
-----
    python bottleneck_detection.py <sensitivity_report.xlsx> [output_report.xlsx]
    python bottleneck_detection.py <sensitivity_report.xlsx> --concepts deer_coat,deer_legs
    python bottleneck_detection.py <sensitivity_report.xlsx> --stage before
    python bottleneck_detection.py <legacy_wide_format.csv> results/bottleneck_report.xlsx

Run ``python bottleneck_detection.py --help`` for every available option/threshold.
example  python bottleneck_detection.py /mnt/sdd/biased_models/multiclass_sensitivity/resnet50/sensitivity_report_resnet50_2026_07_28_21_38_52.xlsx  /mnt/sdd/biased_models/multiclass_sensitivity/resnet50/sensitivity_report_resnet50_cspi_report.xlsx
Loading workbook: /mnt/sdd/biased_models/multiclass_sensitivity/resnet50/sensitivity_report_resnet50_2026_07_28_21_38_52.xlsx

"""

import os
import sys
import re
import math
import argparse

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import seaborn as sns
from openpyxl.styles import Font, PatternFill, Alignment


############################################################
# CONSTANTS
############################################################

META_SHEET_NAMES = {"run_info", "summary", "summary_by_layer", "sensitivity_analysis"}
REQUIRED_LONG_COLUMNS = {"layer_name", "stage", "sensitivity_score", "file_path"}
NEEDED_LONG_COLUMNS = ["class_name", "stage", "layer_name", "lambda_align", "sensitivity_score", "file_path"]
_INVALID_SHEET_CHARS = set("[]:*?/\\")


############################################################
# SMALL UTILITIES
############################################################

_NUM_RE = re.compile(r"(\d+)")


def natural_sort_key(s):
    """Sort key that treats embedded digit runs numerically, e.g. 'features.3' before
    'features.10' (a plain alphabetical sort gets that backwards)."""
    return [int(tok) if tok.isdigit() else tok.lower() for tok in _NUM_RE.split(str(s))]


def sanitize_sheet_name(name, used_names):
    """Make `name` a valid, unique (<=31 chars, no [ ] : * ? / \\) Excel worksheet name,
    tracking already-used names case-insensitively in `used_names` (mutated in place)."""
    cleaned = "".join(ch for ch in str(name) if ch not in _INVALID_SHEET_CHARS).strip() or "Sheet"
    cleaned = cleaned[:31]
    base = cleaned
    suffix = 1
    while cleaned.lower() in used_names:
        suffix += 1
        tail = f"_{suffix}"
        cleaned = f"{base[:31 - len(tail)]}{tail}"
    used_names.add(cleaned.lower())
    return cleaned


def _fs_safe(name):
    """Sanitize a string for use as a file/directory name."""
    return re.sub(r"[^A-Za-z0-9_.-]+", "_", str(name))


def _parse_csv_list(s):
    if not s:
        return None
    return [item.strip() for item in str(s).split(",") if item.strip()]


def stage_tag_label(stage, lambda_val=None):
    """Return (short_tag_for_sheet/file_names, human_readable_label)."""
    if stage == "before":
        return "before", "Before"
    if stage == "legacy":
        return "legacy", "Legacy"
    if lambda_val is None or (isinstance(lambda_val, float) and math.isnan(lambda_val)):
        return "after", "After"
    return f"afterL{int(round(float(lambda_val) * 100)):02d}", f"After(lambda={lambda_val:g})"


############################################################
# EXCEL WORKBOOK DISCOVERY / LOADING
############################################################

def _strip_part_suffix(sheet_name):
    """Undo the '<concept>_p1', '<concept>_p2', ... split used for concepts whose raw row
    count exceeds Excel's per-sheet row limit."""
    return re.sub(r"_p\d+$", "", str(sheet_name))


def load_run_info(xls):
    """Parse the 'Run_Info' key/value sheet (model name, layer execution order, etc.)."""
    if "Run_Info" not in xls.sheet_names:
        return {}
    try:
        df = xls.parse("Run_Info")
    except Exception:
        return {}
    if df.shape[1] < 2:
        return {}
    key_col, val_col = df.columns[0], df.columns[1]
    return {str(k): v for k, v in zip(df[key_col], df[val_col]) if pd.notna(k)}


def build_summary_lookup(xls):
    """concept_name -> dict(Summary sheet row) for TCAV-score lookups, etc."""
    if "Summary" not in xls.sheet_names:
        return {}
    try:
        df = xls.parse("Summary")
    except Exception:
        return {}
    if "concept_name" not in df.columns:
        return {}
    return {row["concept_name"]: row.to_dict() for _, row in df.iterrows()}


def resolve_concept_sheet_map(xls):
    """concept_name -> [sheet names]. Prefers the 'Summary' sheet's authoritative
    'excel_sheet_names' column; falls back to auto-detecting sheets whose header matches
    the expected long-format schema (and merges '<concept>_p1'/'_p2'/... parts back)."""
    sheet_names = xls.sheet_names
    if "Summary" in sheet_names:
        try:
            summary_df = xls.parse("Summary")
        except Exception:
            summary_df = None
        if summary_df is not None and {"concept_name", "excel_sheet_names"}.issubset(summary_df.columns):
            mapping = {}
            for _, row in summary_df.iterrows():
                concept = row["concept_name"]
                if pd.isna(concept):
                    continue
                sheets = [s.strip() for s in str(row["excel_sheet_names"]).split(",") if s.strip()]
                sheets = [s for s in sheets if s in sheet_names]
                mapping[concept] = sheets or [concept]
            if mapping:
                return mapping

    mapping = {}
    for sn in sheet_names:
        if sn.strip().lower() in META_SHEET_NAMES:
            continue
        try:
            head = xls.parse(sn, nrows=1)
        except Exception:
            continue
        if REQUIRED_LONG_COLUMNS.issubset(set(head.columns)):
            mapping.setdefault(_strip_part_suffix(sn), []).append(sn)
    return mapping


def load_concept_long_df(xls, sheet_list):
    """Load (and concatenate, if split across parts) one concept's long-format rows,
    restricted to just the columns this script actually needs."""
    frames = []
    for sn in sheet_list:
        head = xls.parse(sn, nrows=1)
        available = set(head.columns)
        missing_required = REQUIRED_LONG_COLUMNS - available
        if missing_required:
            raise ValueError(f"sheet '{sn}' is missing required column(s): {sorted(missing_required)}")
        usecols = [c for c in NEEDED_LONG_COLUMNS if c in available]
        frames.append(xls.parse(sn, usecols=usecols))
    df = pd.concat(frames, ignore_index=True) if len(frames) > 1 else frames[0]
    if "lambda_align" not in df.columns:
        df["lambda_align"] = np.nan
    return df


def resolve_layer_order(run_info, layers_present):
    """Canonical layer order: prefer Run_Info's 'layers_to_process' (the real execution
    order the report was generated with - alphabetical sort is NOT safe in general, e.g.
    'features.10' sorts before 'features.3'), falling back to a natural sort."""
    layers_present = list(dict.fromkeys(layers_present))
    order_str = (run_info or {}).get("layers_to_process")
    if order_str:
        canonical = [s.strip() for s in str(order_str).split(",") if s.strip()]
        present_set = set(layers_present)
        ordered = [l for l in canonical if l in present_set]
        leftover = [l for l in layers_present if l not in set(ordered)]
        if leftover:
            ordered.extend(sorted(leftover, key=natural_sort_key))
        if ordered:
            return ordered
    return sorted(layers_present, key=natural_sort_key)


def lookup_tcav_score(summary_lookup, concept, stage, lambda_val):
    info = summary_lookup.get(concept)
    if not info:
        return None
    if stage == "before":
        return info.get("tcav_score_before", info.get("tcav_score"))
    if lambda_val is not None and not (isinstance(lambda_val, float) and math.isnan(lambda_val)):
        for key, val in info.items():
            if isinstance(key, str) and key.startswith("tcav_score_after_lambda_"):
                try:
                    key_val = float(key.rsplit("_", 1)[-1])
                except ValueError:
                    continue
                if np.isclose(key_val, float(lambda_val)):
                    return val
    return info.get("tcav_score_after_overall")


def pivot_stage_matrix(concept_df, stage, lambda_val=None, layer_order=None):
    """Long (sample, layer, stage, score) rows -> wide (sample x layer) sensitivity matrix
    for one stage (and, for 'after', one lambda_align). Returns
    (wide_df, dropped_row_count, zero_variance_layers_dropped) or None if nothing matches."""
    sub = concept_df[concept_df["stage"] == stage]
    if stage == "after" and lambda_val is not None:
        sub = sub[np.isclose(sub["lambda_align"].astype(float), float(lambda_val))]
    if sub.empty:
        return None

    wide = sub.pivot_table(index="file_path", columns="layer_name",
                            values="sensitivity_score", aggfunc="mean")
    order = [c for c in (layer_order or []) if c in wide.columns]
    order.extend(sorted([c for c in wide.columns if c not in order], key=natural_sort_key))
    wide = wide[order]
    wide.columns = [str(c).replace("features.", "Layer.") for c in wide.columns]

    n_before = len(wide)
    wide = wide.dropna(axis=0, how="any")
    dropped_rows = n_before - len(wide)
    if wide.empty:
        return wide, dropped_rows, []

    variances = wide.var(axis=0, ddof=0)
    zero_var_cols = variances[variances <= 1e-15].index.tolist()
    if zero_var_cols:
        wide = wide.drop(columns=zero_var_cols)
    return wide, dropped_rows, zero_var_cols


############################################################
# LEGACY WIDE-FORMAT CSV LOADER (kept for backward compatibility)
############################################################

def load_and_prepare_data(csv_file):
    df = pd.read_csv(csv_file)
    try:
        df = df.drop(columns=["Full filepath"])
    except KeyError:
        pass
    try:
        df = df.drop(columns=[c for c in df.columns if c.startswith("sensitivityscore_After_")])
    except KeyError:
        pass
    class_groups = dict(tuple(df.groupby("Full Class Index")))
    for k in class_groups:
        g = class_groups[k].drop(columns=["Full Class Index"])
        g.columns = [c.replace("sensitivityscore_before_", "") for c in g.columns]
        g.columns = [c.replace("features.", "Layer.") for c in g.columns]
        class_groups[k] = g.select_dtypes(include=[np.number])
    return class_groups


############################################################
# GRADIENT ALIGNMENT (uncentered "cosine-like" correlation)
############################################################

def compute_gradient_alignment(df):
    """Gradient-alignment matrix: equivalent to an uncentered correlation, written
    explicitly. Used as the basis for CSPI_M / CSPI_D."""
    S = df.values
    N = S.shape[0]
    numerator = (S.T @ S) / N
    second_moment = np.mean(S ** 2, axis=0)
    denom = np.sqrt(np.outer(second_moment, second_moment))
    alignment = numerator / denom
    return pd.DataFrame(alignment, index=df.columns, columns=df.columns)


############################################################
# CONCEPT SENSITIVITY PROPAGATION INDEX (CSPI_M / CSPI_D)
############################################################

def compute_cspi(alignment_matrix, tau_abs=0.1, tau_rel=0.1):
    """
    CSPI_M ("magnitude"): mean |alignment| of a layer with every downstream layer - how
        strongly the concept signal at this layer keeps showing up later, regardless of sign.
    CSPI_D ("directional"): mean *signed* alignment with every downstream layer - positive
        means the concept keeps pointing the same way; negative flags a sign flip.

    Layer_Type / Is_Bottleneck labeling:
        CSPI_M < tau_abs               -> "Weak"              (Is_Bottleneck = 0)
        CSPI_D > tau_rel               -> "Positive_Aligned"   (Is_Bottleneck = 0)
        CSPI_D < -tau_rel              -> "Negative_Aligned"   (Is_Bottleneck = 1, sign flip)
        otherwise                      -> "Mixed_Strong"       (Is_Bottleneck = 1, strong but unclear)
    """
    layers = alignment_matrix.columns
    n = len(layers)
    scores_absolute, scores_relative, layer_type, is_bottleneck = [], [], [], []

    for i in range(n):
        downstream = alignment_matrix.iloc[i, i + 1:]
        if len(downstream) > 0:
            score_abs = float(np.mean(np.abs(downstream)))
            score_rel = float(np.mean(downstream))
        else:
            score_abs, score_rel = 0.0, 0.0

        if score_abs < tau_abs:
            label, bottleneck = "Weak", 0
        elif score_rel > tau_rel:
            label, bottleneck = "Positive_Aligned", 0
        elif score_rel < -tau_rel:
            label, bottleneck = "Negative_Aligned", 1
        else:
            label, bottleneck = "Mixed_Strong", 1

        scores_absolute.append(score_abs)
        scores_relative.append(score_rel)
        layer_type.append(label)
        is_bottleneck.append(bottleneck)

    result = pd.DataFrame({
        "Layer": layers,
        "CSPI_M": scores_absolute,
        "CSPI_D": scores_relative,
        "Layer_Type": layer_type,
        "Is_Bottleneck": is_bottleneck,
    })
    result["Bottleneck_Rank"] = result["CSPI_M"].rank(ascending=False, method="dense").astype(int)
    return result


############################################################
# ADJACENT-LAYER CORRELATION / DELTA / MOVING AVERAGE
############################################################

def compute_adjacent_correlation(corr_matrix):
    """R(i, i+1) for every consecutive layer pair; length == n_layers - 1."""
    n = corr_matrix.shape[0]
    return np.array([corr_matrix[i, i + 1] for i in range(n - 1)])


def compute_delta(R, column_names):
    """Change in adjacent correlation from one pair to the next; length == n_layers - 2."""
    delta = np.diff(R)
    layer_pairs = [(column_names[i], column_names[i + 1]) for i in range(len(delta))]
    return delta, layer_pairs


def compute_moving_average(deltas, labels, window_size=3):
    """Moving average of `deltas`, paired 1:1 with `labels` (must be the same length -
    e.g. the layer-pair labels returned by compute_delta)."""
    deltas = np.asarray(deltas, dtype=float)
    labels = list(labels)
    n = len(deltas)
    if n == 0:
        return pd.DataFrame({"Moving_Avg_Delta": [], "Window_Layers": []})
    if n < window_size:
        return pd.DataFrame({"Moving_Avg_Delta": deltas, "Window_Layers": labels})

    moving_avg = np.convolve(deltas, np.ones(window_size) / window_size, mode="valid")
    pad_size = n - len(moving_avg)
    moving_avg = np.concatenate([deltas[:pad_size], moving_avg])

    window_layers = []
    for i in range(n):
        if i < pad_size:
            window_layers.append(labels[i])
        else:
            w0 = i - pad_size
            window_layers.append(" | ".join(labels[w0:w0 + window_size]))
    return pd.DataFrame({"Moving_Avg_Delta": moving_avg, "Window_Layers": window_layers})


############################################################
# BOTTLENECK DETECTION
############################################################

def detect_bottlenecks_adjacent(R, delta, cspi_m=None, tau_delta=0.2, tau_corr=0.2, tau_cspi=0.6):
    """
    Adjacent-layer-PAIR bottleneck flags (one entry per pair i -> i+1; length == len(R)).
    A pair is flagged when its correlation sharply drops / is weak / flips sign, AND (if
    cspi_m is given) the downstream CSPI_M gate at layer i is still below tau_cspi.
    """
    flags = np.zeros(len(R), dtype=int)
    reasons = [""] * len(R)
    for i in range(len(R)):
        cond_drop = i < len(delta) and delta[i] < -tau_delta
        cond_weak = R[i] < tau_corr
        cond_flip = R[i] < 0
        cond_gate = True if cspi_m is None else (i < len(cspi_m) and cspi_m[i] < tau_cspi)
        triggered = (cond_drop or cond_weak or cond_flip) and cond_gate
        flags[i] = int(triggered)
        if triggered:
            parts = []
            if cond_flip:
                parts.append("sign flip")
            elif cond_weak:
                parts.append("weak correlation")
            if cond_drop:
                parts.append("sharp drop")
            reasons[i] = ", ".join(parts) if parts else "flagged"
    return flags, reasons


def detect_bottlenecks_threshold(alignment_matrix, threshold=0.2):
    """Legacy/simple detector kept for backward compatibility: flags a layer whenever its
    mean |alignment| with every downstream layer is >= threshold (no directionality)."""
    layers = alignment_matrix.columns
    bottlenecks = []
    for i in range(len(layers)):
        downstream = alignment_matrix.iloc[i, i + 1:]
        if len(downstream) == 0:
            continue
        mean_corr = float(np.mean(np.abs(downstream)))
        if mean_corr >= threshold:
            bottlenecks.append({"Layer": layers[i], "Mean_Alignment": mean_corr})
    return pd.DataFrame(bottlenecks)


############################################################
# TOP-K CORRELATION SUBSET
############################################################

def compute_top_subset(corr_matrix, top_k=10):
    last_row = corr_matrix.iloc[-1]
    positive = last_row[last_row > 0]
    if len(positive) <= 1:
        return None
    top_cols = positive.sort_values(ascending=False).head(top_k).index.tolist()
    return corr_matrix.loc[top_cols, top_cols]


############################################################
# UNIFIED PER-LAYER TABLE (one concept + one stage -> one table)
############################################################

def build_layer_table(wide_df, tau_abs, tau_rel, tau_delta, tau_corr, tau_cspi, window_size):
    corr = wide_df.corr()
    alignment = compute_gradient_alignment(wide_df)
    cspi_df = compute_cspi(alignment, tau_abs=tau_abs, tau_rel=tau_rel)

    layers = wide_df.columns.tolist()
    n = len(layers)
    R = compute_adjacent_correlation(corr.values)
    delta, layer_pairs = compute_delta(R, layers)
    pair_labels = [f"{a} -> {b}" for a, b in layer_pairs]
    moving_avg_df = compute_moving_average(delta, pair_labels, window_size=window_size)

    adj_flags, adj_reasons = detect_bottlenecks_adjacent(
        R, delta, cspi_m=cspi_df["CSPI_M"].values,
        tau_delta=tau_delta, tau_corr=tau_corr, tau_cspi=tau_cspi)

    table = cspi_df.copy()
    table.insert(1, "Layer_Index", np.arange(n))

    adjacent_r_full = np.full(n, np.nan)
    adjacent_r_full[:len(R)] = R
    table["Adjacent_R"] = adjacent_r_full

    # Aligned to `delta[i]` (== R[i+1]-R[i]), the SAME value detect_bottlenecks_adjacent()
    # used to decide row i's flag - NOT a plain .diff() of Adjacent_R, which would land the
    # value one row later than the flag/reason it explains.
    delta_full = np.full(n, np.nan)
    delta_full[:len(delta)] = delta
    table["Delta_R"] = delta_full

    adj_flag_full = np.zeros(n, dtype=int)
    adj_flag_full[:len(adj_flags)] = adj_flags
    table["Is_Bottleneck_AdjacentDrop"] = adj_flag_full
    table["Adjacent_Bottleneck_Reason"] = list(adj_reasons) + [""] * (n - len(adj_reasons))

    table["Is_Bottleneck_Any"] = ((table["Is_Bottleneck"] == 1) |
                                   (table["Is_Bottleneck_AdjacentDrop"] == 1)).astype(int)

    n_pairs = len(pair_labels)
    delta_detail = pd.DataFrame({
        "Layer_Pair": pair_labels,
        "R_i": R[:n_pairs],
        "R_i+1": R[1:1 + n_pairs],
        "Delta_R": delta,
        "Moving_Avg_Delta": (moving_avg_df["Moving_Avg_Delta"].values
                              if len(moving_avg_df) == n_pairs else np.full(n_pairs, np.nan)),
    })

    return {"table": table, "corr": corr, "alignment": alignment, "delta_df": delta_detail}


############################################################
# BEFORE / AFTER COMPARISON
############################################################

def build_comparison_table(table_before, table_after, before_label, after_label):
    b = table_before.set_index("Layer")
    a = table_after.set_index("Layer")
    common = [l for l in b.index if l in a.index]
    if not common:
        return pd.DataFrame()

    cmp = pd.DataFrame({"Layer": common})
    for metric in ("CSPI_M", "CSPI_D", "Adjacent_R"):
        cmp[f"{metric}_{before_label}"] = b.loc[common, metric].values
        cmp[f"{metric}_{after_label}"] = a.loc[common, metric].values
        cmp[f"Delta_{metric}"] = cmp[f"{metric}_{after_label}"] - cmp[f"{metric}_{before_label}"]

    cmp[f"Bottleneck_{before_label}"] = b.loc[common, "Is_Bottleneck_Any"].values
    cmp[f"Bottleneck_{after_label}"] = a.loc[common, "Is_Bottleneck_Any"].values

    def _status(row):
        bb, aa = row[f"Bottleneck_{before_label}"], row[f"Bottleneck_{after_label}"]
        if bb and aa:
            return "Persisting"
        if bb and not aa:
            return "Resolved"
        if (not bb) and aa:
            return "Introduced"
        return "None"

    cmp["Bottleneck_Status"] = cmp.apply(_status, axis=1)
    return cmp


############################################################
# PLOTTING
############################################################

def plot_heatmap(matrix, title, save_path, annot=None):
    n = matrix.shape[0]
    if annot is None:
        annot = n <= 20
    mask = np.triu(np.ones_like(matrix, dtype=bool))
    side = max(6.0, min(22.0, 0.32 * n + 3))
    plt.figure(figsize=(side, side * 0.85))
    sns.heatmap(matrix, annot=annot, fmt=".2f" if annot else "", cmap="coolwarm",
                mask=mask, square=True, cbar_kws={"shrink": 0.8}, vmin=-1, vmax=1, center=0)
    plt.title(title)
    tick_fs = max(5, 9 - n // 12)
    plt.xticks(rotation=90 if n > 12 else 45, fontsize=tick_fs)
    plt.yticks(rotation=0, fontsize=tick_fs)
    plt.tight_layout()
    plt.savefig(save_path, dpi=130)
    plt.close()


def plot_cspi_bar(table, title, save_path):
    n = len(table)
    fig_w = max(10, 0.28 * n)
    fig, ax = plt.subplots(figsize=(fig_w, 6))
    x = np.arange(n)
    ax.bar(x - 0.2, table["CSPI_M"], width=0.4, label="CSPI_M (magnitude)", color="#4472C4")
    ax.bar(x + 0.2, table["CSPI_D"], width=0.4, label="CSPI_D (directional)", color="#ED7D31")
    for pos in np.where(table["Is_Bottleneck_Any"].values == 1)[0]:
        ax.axvspan(pos - 0.5, pos + 0.5, color="red", alpha=0.12, zorder=0)
    ax.axhline(0, color="black", linewidth=0.8)
    ax.set_xticks(x)
    ax.set_xticklabels(table["Layer"], rotation=90, fontsize=max(5, 9 - n // 15))
    ax.set_ylabel("CSPI score")
    ax.set_title(title + "\n(red shading = flagged bottleneck layer)")
    ax.legend(loc="upper right")
    fig.tight_layout()
    fig.savefig(save_path, dpi=130)
    plt.close(fig)


def plot_adjacent_line(table, title, save_path):
    n = len(table)
    fig_w = max(10, 0.28 * n)
    fig, ax = plt.subplots(figsize=(fig_w, 6))
    x = np.arange(n)
    ax.plot(x, table["Adjacent_R"], marker="o", markersize=3, color="#2E75B6",
            label="Adjacent R(i, i+1)")
    ax.axhline(0, color="black", linewidth=0.8)
    flagged = np.where(table["Is_Bottleneck_AdjacentDrop"].values == 1)[0]
    if len(flagged):
        ax.scatter(flagged, table["Adjacent_R"].values[flagged], color="red", zorder=5,
                   label="Adjacent-drop bottleneck")
    ax.set_xticks(x)
    ax.set_xticklabels(table["Layer"], rotation=90, fontsize=max(5, 9 - n // 15))
    ax.set_ylabel("Adjacent correlation R(i, i+1)")
    ax.set_title(title)
    ax.legend(loc="upper right")
    fig.tight_layout()
    fig.savefig(save_path, dpi=130)
    plt.close(fig)


def plot_comparison_bar(cmp_table, before_label, after_label, title, save_path):
    n = len(cmp_table)
    fig_w = max(10, 0.28 * n)
    fig, ax = plt.subplots(figsize=(fig_w, 6))
    x = np.arange(n)
    ax.bar(x - 0.2, cmp_table[f"CSPI_M_{before_label}"], width=0.4,
           label=f"CSPI_M {before_label}", color="#4472C4")
    ax.bar(x + 0.2, cmp_table[f"CSPI_M_{after_label}"], width=0.4,
           label=f"CSPI_M {after_label}", color="#70AD47")
    for pos in np.where(cmp_table["Bottleneck_Status"].values == "Introduced")[0]:
        ax.axvspan(pos - 0.5, pos + 0.5, color="red", alpha=0.12)
    for pos in np.where(cmp_table["Bottleneck_Status"].values == "Resolved")[0]:
        ax.axvspan(pos - 0.5, pos + 0.5, color="green", alpha=0.12)
    ax.axhline(0, color="black", linewidth=0.8)
    ax.set_xticks(x)
    ax.set_xticklabels(cmp_table["Layer"], rotation=90, fontsize=max(5, 9 - n // 15))
    ax.set_ylabel("CSPI_M")
    ax.set_title(title + "\n(green = bottleneck resolved, red = bottleneck introduced)")
    ax.legend(loc="upper right")
    fig.tight_layout()
    fig.savefig(save_path, dpi=130)
    plt.close(fig)


def plot_comparison_adjacent(cmp_table, before_label, after_label, title, save_path):
    n = len(cmp_table)
    fig_w = max(10, 0.28 * n)
    fig, ax = plt.subplots(figsize=(fig_w, 6))
    x = np.arange(n)
    ax.plot(x, cmp_table[f"Adjacent_R_{before_label}"], marker="o", markersize=3,
            label=f"Adjacent R {before_label}", color="#4472C4")
    ax.plot(x, cmp_table[f"Adjacent_R_{after_label}"], marker="s", markersize=3,
            label=f"Adjacent R {after_label}", color="#70AD47")
    ax.axhline(0, color="black", linewidth=0.8)
    ax.set_xticks(x)
    ax.set_xticklabels(cmp_table["Layer"], rotation=90, fontsize=max(5, 9 - n // 15))
    ax.set_ylabel("Adjacent correlation R(i, i+1)")
    ax.set_title(title)
    ax.legend(loc="upper right")
    fig.tight_layout()
    fig.savefig(save_path, dpi=130)
    plt.close(fig)


def plot_overview_bar(overview_df, save_path):
    pivot = overview_df.pivot_table(index="Concept", columns="Stage", values="Mean_CSPI_M", aggfunc="mean")
    ax = pivot.plot(kind="bar", figsize=(max(10, 0.5 * len(pivot)), 6))
    ax.set_ylabel("Mean CSPI_M across layers")
    ax.set_title("Mean concept-sensitivity propagation (CSPI_M) by concept")
    ax.legend(title="Stage")
    plt.xticks(rotation=45, ha="right")
    fig = ax.get_figure()
    fig.tight_layout()
    fig.savefig(save_path, dpi=130)
    plt.close(fig)


def plot_overview_tcav(overview_df, save_path):
    if "TCAV_Score" not in overview_df.columns:
        return None
    sub = overview_df.dropna(subset=["TCAV_Score"])
    if sub.empty:
        return None
    pivot = sub.pivot_table(index="Concept", columns="Stage", values="TCAV_Score", aggfunc="mean")
    ax = pivot.plot(kind="bar", figsize=(max(10, 0.5 * len(pivot)), 6))
    ax.axhline(0.5, color="black", linewidth=0.8, linestyle="--")
    ax.set_ylabel("TCAV score")
    ax.set_title("TCAV score by concept (from the Summary sheet)")
    ax.legend(title="Stage")
    plt.xticks(rotation=45, ha="right")
    fig = ax.get_figure()
    fig.tight_layout()
    fig.savefig(save_path, dpi=130)
    plt.close(fig)
    return save_path


############################################################
# EXCEL WRITING HELPERS
############################################################

def _embed_image(worksheet, image_path, anchor_cell, max_width_px=780):
    if not image_path or not os.path.isfile(image_path):
        return
    try:
        from openpyxl.drawing.image import Image as XLImage
        img = XLImage(image_path)
        if img.width and img.width > max_width_px:
            scale = max_width_px / float(img.width)
            img.width = int(img.width * scale)
            img.height = int(img.height * scale)
        worksheet.add_image(img, anchor_cell)
    except Exception as e:
        print(f"  (warning) could not embed image '{image_path}': {e}")


def _round_numeric(df, decimals=4):
    return df.round(decimals) if df is not None and not df.empty else df


def write_df_sheet(writer, df, base_name, used_sheet_names, index=False):
    sheet_name = sanitize_sheet_name(base_name, used_sheet_names)
    df.to_excel(writer, sheet_name=sheet_name, index=index)
    ws = writer.sheets[sheet_name]
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    try:
        ws.auto_filter.ref = ws.dimensions
    except Exception:
        pass
    for col_cells in ws.columns:
        max_len = 0
        col_letter = col_cells[0].column_letter
        for cell in col_cells:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 45)
    return sheet_name, ws


############################################################
# PER-CONCEPT / PER-STAGE PROCESSING (xlsx pipeline)
############################################################

def process_concept_stage(concept, long_df, stage, lambda_val, layer_order_all,
                           args, thresholds, charts_dir, summary_lookup):
    tag, label = stage_tag_label(stage, lambda_val)
    pivoted = pivot_stage_matrix(long_df, stage, lambda_val, layer_order_all)
    if pivoted is None:
        print(f"  [{concept}/{tag}] no rows for this stage - skipped")
        return None
    wide, dropped_rows, dropped_layers = pivoted
    if wide.shape[1] < 2 or wide.shape[0] < args.min_samples:
        print(f"  [{concept}/{tag}] not enough usable data after cleaning "
              f"({wide.shape[0]} samples x {wide.shape[1]} layers) - skipped")
        return None

    built = build_layer_table(wide, thresholds["tau_abs"], thresholds["tau_rel"],
                               thresholds["tau_delta"], thresholds["tau_corr"],
                               thresholds["tau_cspi"], args.window_size)
    table = built["table"]
    n_bottleneck = int(table["Is_Bottleneck_Any"].sum())
    msg = (f"  [{concept}/{tag}] {wide.shape[0]} samples x {wide.shape[1]} layers "
           f"-> {n_bottleneck} bottleneck layer(s) flagged")
    if dropped_rows:
        msg += f" ({dropped_rows} sample(s) dropped for missing layers)"
    if dropped_layers:
        msg += f" ({len(dropped_layers)} zero-variance layer(s) excluded)"
    print(msg)

    charts = {}
    if not args.no_charts:
        heatmap_path = os.path.join(charts_dir, f"{tag}_heatmap.png")
        plot_heatmap(built["corr"], f"{concept} [{label}] - layer correlation", heatmap_path)
        charts["heatmap"] = heatmap_path

        subset = compute_top_subset(built["corr"], top_k=args.top_k)
        if subset is not None:
            subset_path = os.path.join(charts_dir, f"{tag}_subset_heatmap.png")
            plot_heatmap(subset, f"{concept} [{label}] - top-{args.top_k} correlated layers",
                         subset_path, annot=True)
            charts["subset"] = subset_path

        bar_path = os.path.join(charts_dir, f"{tag}_cspi_bar.png")
        plot_cspi_bar(table, f"{concept} [{label}] - CSPI by layer", bar_path)
        charts["cspi_bar"] = bar_path

        line_path = os.path.join(charts_dir, f"{tag}_adjacent_r.png")
        plot_adjacent_line(table, f"{concept} [{label}] - adjacent-layer correlation", line_path)
        charts["adjacent_line"] = line_path

    return {
        "stage": stage, "lambda_val": lambda_val, "tag": tag, "label": label,
        "table": table, "delta_df": built["delta_df"], "corr": built["corr"],
        "n_samples": int(wide.shape[0]), "n_layers": int(wide.shape[1]),
        "dropped_rows": int(dropped_rows), "dropped_layers": dropped_layers,
        "charts": charts,
        "tcav_score": lookup_tcav_score(summary_lookup, concept, stage, lambda_val),
    }


def process_concept(concept, xls, sheet_list, run_info, summary_lookup, args, thresholds,
                     charts_root, lambdas_wanted):
    long_df = load_concept_long_df(xls, sheet_list)
    class_name = long_df["class_name"].iloc[0] if "class_name" in long_df.columns and len(long_df) else None
    stages_present = set(long_df["stage"].dropna().unique().tolist())

    want_before = args.stage in ("auto", "before") and "before" in stages_present
    want_after = args.stage in ("auto", "after") and "after" in stages_present

    layer_order_all = resolve_layer_order(run_info, long_df["layer_name"].dropna().unique().tolist())
    concept_charts_dir = os.path.join(charts_root, _fs_safe(concept))
    os.makedirs(concept_charts_dir, exist_ok=True)

    stage_results = {}

    if want_before:
        r = process_concept_stage(concept, long_df, "before", None, layer_order_all,
                                   args, thresholds, concept_charts_dir, summary_lookup)
        if r:
            stage_results["before"] = r

    if want_after:
        lambdas_present = sorted(long_df.loc[long_df["stage"] == "after", "lambda_align"].dropna().unique().tolist())
        if lambdas_wanted:
            lambdas_present = [l for l in lambdas_present if any(np.isclose(l, w) for w in lambdas_wanted)]
        for lam in lambdas_present:
            tag, _ = stage_tag_label("after", lam)
            r = process_concept_stage(concept, long_df, "after", lam, layer_order_all,
                                       args, thresholds, concept_charts_dir, summary_lookup)
            if r:
                stage_results[tag] = r

    comparisons = {}
    if "before" in stage_results:
        for tag, r in stage_results.items():
            if tag == "before":
                continue
            cmp_table = build_comparison_table(stage_results["before"]["table"], r["table"],
                                                "Before", r["label"])
            if cmp_table.empty:
                continue
            cmp_charts = {}
            if not args.no_charts:
                bar_path = os.path.join(concept_charts_dir, f"{tag}_vs_before_bar.png")
                plot_comparison_bar(cmp_table, "Before", r["label"],
                                     f"{concept}: CSPI_M Before vs {r['label']}", bar_path)
                line_path = os.path.join(concept_charts_dir, f"{tag}_vs_before_adjacent.png")
                plot_comparison_adjacent(cmp_table, "Before", r["label"],
                                          f"{concept}: Adjacent R Before vs {r['label']}", line_path)
                cmp_charts = {"bar": bar_path, "line": line_path}
            comparisons[tag] = {"before_label": "Before", "after_label": r["label"],
                                 "cmp_table": cmp_table, "charts": cmp_charts}

    if not stage_results:
        print(f"  (skipped '{concept}': no requested stage had usable data)")
        return None

    return {"class_name": class_name, "stages": stage_results, "comparisons": comparisons}


def run_xlsx_pipeline(input_path, args, thresholds, charts_root):
    print(f"Loading workbook: {input_path}")
    xls = pd.ExcelFile(input_path, engine="openpyxl")
    run_info = load_run_info(xls)
    summary_lookup = build_summary_lookup(xls)
    concept_sheet_map = resolve_concept_sheet_map(xls)
    if not concept_sheet_map:
        raise ValueError("No concept sheets were found in this workbook (expected long-format "
                          "sheets with columns like layer_name/stage/sensitivity_score/file_path, "
                          "or a 'Summary' sheet listing them under 'excel_sheet_names').")

    concepts_wanted = _parse_csv_list(args.concepts)
    if concepts_wanted:
        wanted_lower = {c.lower() for c in concepts_wanted}
        concept_sheet_map = {k: v for k, v in concept_sheet_map.items() if k.lower() in wanted_lower}
        missing = wanted_lower - {k.lower() for k in concept_sheet_map}
        if missing:
            print(f"WARNING: requested concept(s) not found in workbook: {sorted(missing)}")

    lambdas_wanted = _parse_csv_list(args.lambdas)
    lambdas_wanted = {float(v) for v in lambdas_wanted} if lambdas_wanted else None

    all_results = {}
    n_concepts = len(concept_sheet_map)
    for idx, (concept, sheet_list) in enumerate(concept_sheet_map.items(), start=1):
        print(f"[{idx}/{n_concepts}] Processing concept '{concept}' (sheet(s): {sheet_list}) ...")
        try:
            result = process_concept(concept, xls, sheet_list, run_info, summary_lookup,
                                      args, thresholds, charts_root, lambdas_wanted)
        except Exception as e:
            print(f"  ERROR processing concept '{concept}': {e}")
            continue
        if result:
            all_results[concept] = result
    return run_info, all_results


############################################################
# LEGACY CSV PIPELINE
############################################################

def run_csv_pipeline(csv_file, args, thresholds, charts_root):
    class_groups = load_and_prepare_data(csv_file)
    all_results = {}
    n_classes = len(class_groups)
    for idx, (class_id, df) in enumerate(class_groups.items(), start=1):
        concept = f"Class_{class_id}"
        print(f"[{idx}/{n_classes}] Processing legacy CSV class '{class_id}' ...")

        variances = df.var(axis=0, ddof=0)
        zero_var_cols = variances[variances <= 1e-15].index.tolist()
        if zero_var_cols:
            df = df.drop(columns=zero_var_cols)
        if df.shape[1] < 2 or df.shape[0] < args.min_samples:
            print(f"  (skipped '{concept}': not enough numeric layer columns/samples)")
            continue

        built = build_layer_table(df, thresholds["tau_abs"], thresholds["tau_rel"],
                                   thresholds["tau_delta"], thresholds["tau_corr"],
                                   thresholds["tau_cspi"], args.window_size)
        table = built["table"]
        n_bottleneck = int(table["Is_Bottleneck_Any"].sum())
        print(f"  [{concept}] {df.shape[0]} samples x {df.shape[1]} layers "
              f"-> {n_bottleneck} bottleneck layer(s) flagged")

        concept_charts_dir = os.path.join(charts_root, _fs_safe(concept))
        os.makedirs(concept_charts_dir, exist_ok=True)
        charts = {}
        if not args.no_charts:
            heatmap_path = os.path.join(concept_charts_dir, "legacy_heatmap.png")
            plot_heatmap(built["corr"], f"{concept} - layer correlation", heatmap_path)
            charts["heatmap"] = heatmap_path
            subset = compute_top_subset(built["corr"], top_k=args.top_k)
            if subset is not None:
                subset_path = os.path.join(concept_charts_dir, "legacy_subset_heatmap.png")
                plot_heatmap(subset, f"{concept} - top-{args.top_k} correlated layers", subset_path, annot=True)
                charts["subset"] = subset_path
            bar_path = os.path.join(concept_charts_dir, "legacy_cspi_bar.png")
            plot_cspi_bar(table, f"{concept} - CSPI by layer", bar_path)
            charts["cspi_bar"] = bar_path
            line_path = os.path.join(concept_charts_dir, "legacy_adjacent_r.png")
            plot_adjacent_line(table, f"{concept} - adjacent-layer correlation", line_path)
            charts["adjacent_line"] = line_path

        all_results[concept] = {
            "class_name": concept,
            "stages": {"legacy": {
                "stage": "legacy", "lambda_val": None, "tag": "legacy", "label": "Legacy",
                "table": table, "delta_df": built["delta_df"], "corr": built["corr"],
                "n_samples": int(df.shape[0]), "n_layers": int(df.shape[1]),
                "dropped_rows": 0, "dropped_layers": zero_var_cols, "charts": charts,
                "tcav_score": None,
            }},
            "comparisons": {},
        }
    return all_results


############################################################
# WORKBOOK ASSEMBLY (single writer pass, controlled sheet order)
############################################################

def write_workbook(output_xlsx, all_results, run_info, args):
    used_sheet_names = set()
    bottleneck_rows = []
    comparison_rows = []
    overview_rows = []

    for concept, result in all_results.items():
        class_name = result.get("class_name")
        for tag, r in result["stages"].items():
            table = r["table"]
            n_bottleneck = int(table["Is_Bottleneck_Any"].sum())
            overview_rows.append({
                "Concept": concept, "Class": class_name, "Stage": r["label"],
                "N_Layers": r["n_layers"], "N_Samples": r["n_samples"],
                "Mean_CSPI_M": table["CSPI_M"].mean(), "Mean_CSPI_D": table["CSPI_D"].mean(),
                "N_Bottleneck_Layers": n_bottleneck,
                "Pct_Bottleneck_Layers": round(100.0 * n_bottleneck / max(1, r["n_layers"]), 2),
                "TCAV_Score": r.get("tcav_score"),
            })
            for _, row in table.iterrows():
                if row["Is_Bottleneck_Any"] != 1:
                    continue
                bottleneck_rows.append({
                    "Concept": concept, "Class": class_name, "Stage": r["label"],
                    "Layer_Index": row["Layer_Index"], "Layer": row["Layer"],
                    "CSPI_M": row["CSPI_M"], "CSPI_D": row["CSPI_D"],
                    "Layer_Type": row["Layer_Type"],
                    "Is_Bottleneck_GlobalDownstream": row["Is_Bottleneck"],
                    "Is_Bottleneck_AdjacentDrop": row["Is_Bottleneck_AdjacentDrop"],
                    "Bottleneck_Rank": row["Bottleneck_Rank"],
                    "Adjacent_R": row["Adjacent_R"], "Delta_R": row["Delta_R"],
                    "Reason": row["Adjacent_Bottleneck_Reason"] or row["Layer_Type"],
                })
        for tag, cmp in result["comparisons"].items():
            for _, row in cmp["cmp_table"].iterrows():
                comparison_rows.append({
                    "Concept": concept, "Class": class_name,
                    "Comparison": f'{cmp["before_label"]} vs {cmp["after_label"]}',
                    **row.to_dict(),
                })

    overview_df = pd.DataFrame(overview_rows)
    bottleneck_df = pd.DataFrame(bottleneck_rows)
    comparison_df = pd.DataFrame(comparison_rows) if comparison_rows else pd.DataFrame()

    result_dir = os.path.dirname(os.path.abspath(output_xlsx)) or "."
    overview_charts = {}
    if not overview_df.empty and not args.no_charts:
        overview_dir = os.path.join(result_dir, "charts")
        os.makedirs(overview_dir, exist_ok=True)
        bar_path = os.path.join(overview_dir, "overview_cspi_by_concept.png")
        plot_overview_bar(overview_df, bar_path)
        overview_charts["cspi_bar"] = bar_path
        tcav_path = os.path.join(overview_dir, "overview_tcav_by_concept.png")
        if plot_overview_tcav(overview_df, tcav_path):
            overview_charts["tcav_bar"] = tcav_path

    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:
        if run_info:
            run_info_df = pd.DataFrame([(k, str(v)) for k, v in run_info.items()],
                                        columns=["Parameter", "Value"])
            write_df_sheet(writer, run_info_df, "Run_Info", used_sheet_names)

        if not overview_df.empty:
            _, ws = write_df_sheet(writer, _round_numeric(overview_df), "Overview", used_sheet_names)
            chart_row = len(overview_df) + 3
            for key in ("cspi_bar", "tcav_bar"):
                if key in overview_charts:
                    _embed_image(ws, overview_charts[key], f"A{chart_row}")
                    chart_row += 32

        if not bottleneck_df.empty:
            bottleneck_df = bottleneck_df.sort_values(
                ["Concept", "Stage", "Bottleneck_Rank"]).reset_index(drop=True)
            write_df_sheet(writer, _round_numeric(bottleneck_df), "Bottleneck_Summary", used_sheet_names)
        else:
            print("(No layer crossed the bottleneck thresholds in any concept/stage - "
                  "'Bottleneck_Summary' sheet omitted.)")

        if not comparison_df.empty:
            write_df_sheet(writer, _round_numeric(comparison_df), "Comparison_Summary", used_sheet_names)

        for concept, result in all_results.items():
            for tag, r in result["stages"].items():
                base = f"{concept}_{tag}"

                _, ws_tbl = write_df_sheet(writer, _round_numeric(r["table"]), f"{base}_tbl", used_sheet_names)
                chart_row = len(r["table"]) + 3
                for key in ("cspi_bar", "adjacent_line"):
                    if key in r["charts"]:
                        _embed_image(ws_tbl, r["charts"][key], f"A{chart_row}")
                        chart_row += 32

                corr_display = r["corr"].mask(np.triu(np.ones_like(r["corr"], dtype=bool)))
                corr_display = corr_display.reset_index().rename(columns={"index": "Layer"})
                _, ws_corr = write_df_sheet(writer, _round_numeric(corr_display), f"{base}_corr", used_sheet_names)
                chart_row = len(corr_display) + 3
                for key in ("heatmap", "subset"):
                    if key in r["charts"]:
                        _embed_image(ws_corr, r["charts"][key], f"A{chart_row}")
                        chart_row += 34

                if r["delta_df"] is not None and not r["delta_df"].empty:
                    write_df_sheet(writer, _round_numeric(r["delta_df"]), f"{base}_delta", used_sheet_names)

            for tag, cmp in result["comparisons"].items():
                base = f"{concept}_{tag}"
                _, ws_cmp = write_df_sheet(writer, _round_numeric(cmp["cmp_table"]), f"{base}_cmp", used_sheet_names)
                chart_row = len(cmp["cmp_table"]) + 3
                for key in ("bar", "line"):
                    if key in cmp["charts"]:
                        _embed_image(ws_cmp, cmp["charts"][key], f"A{chart_row}")
                        chart_row += 32

    return output_xlsx


############################################################
# CLI
############################################################

def default_output_path(input_file):
    base = os.path.splitext(os.path.basename(input_file))[0]
    directory = os.path.dirname(os.path.abspath(input_file))
    return os.path.join(directory, f"bottleneck_report_{base}.xlsx")


def parse_args():
    parser = argparse.ArgumentParser(
        description="Compute CSPI_M / CSPI_D concept-sensitivity-propagation scores and detect "
                    "layer-wise bottlenecks from a (multiclass) sensitivity report, with optional "
                    "before/after recalibration comparison driven by the report's 'stage' column.")
    parser.add_argument("input_file",
                         help="Path to the sensitivity report (.xlsx, from utils_sensitivity_multiclass.py) "
                              "or a legacy wide-format .csv")
    parser.add_argument("output_xlsx", nargs="?", default=None,
                         help="Path of the bottleneck-report .xlsx to write "
                              "(default: alongside the input file)")
    parser.add_argument("--stage", choices=["auto", "before", "after"], default="auto",
                         help="Which stage(s) to analyze. 'auto' (default) uses whatever stage(s) "
                              "are present per concept, and automatically builds a before/after "
                              "comparison whenever both are found.")
    parser.add_argument("--lambdas", default=None,
                         help="Comma-separated lambda_align values to restrict the 'after' stage to "
                              "(default: every lambda found in the data)")
    parser.add_argument("--concepts", default=None,
                         help="Comma-separated concept names to process, e.g. 'deer_coat,deer_legs' "
                              "(default: every concept sheet found)")
    parser.add_argument("--tau-abs", type=float, default=0.1,
                         help="CSPI_M threshold below which a layer's downstream alignment is 'Weak' (default 0.1)")
    parser.add_argument("--tau-rel", type=float, default=0.1,
                         help="CSPI_D threshold for Positive/Negative-aligned labeling (default 0.1)")
    parser.add_argument("--tau-delta", type=float, default=0.2,
                         help="Adjacent-correlation drop threshold for the adjacent-layer bottleneck detector (default 0.2)")
    parser.add_argument("--tau-corr", type=float, default=0.2,
                         help="Adjacent-correlation weak threshold for the adjacent-layer bottleneck detector (default 0.2)")
    parser.add_argument("--tau-cspi", type=float, default=0.6,
                         help="CSPI_M gate used by the adjacent-layer bottleneck detector (default 0.6)")
    parser.add_argument("--top-k", type=int, default=10,
                         help="Number of layers kept in the top-k correlation subset heatmap (default 10)")
    parser.add_argument("--window-size", type=int, default=3,
                         help="Moving-average window size over adjacent-correlation deltas (default 3)")
    parser.add_argument("--min-samples", type=int, default=5,
                         help="Minimum complete samples required to analyze a concept/stage (default 5)")
    parser.add_argument("--no-charts", action="store_true",
                         help="Skip PNG chart generation/embedding (tables only, faster)")
    return parser.parse_args()


def main():
    args = parse_args()
    input_path = args.input_file
    if not os.path.isfile(input_path):
        print(f"ERROR: input file not found: {input_path}")
        sys.exit(1)

    ext = os.path.splitext(input_path)[1].lower()
    output_xlsx = args.output_xlsx or default_output_path(input_path)
    result_dir = os.path.dirname(os.path.abspath(output_xlsx)) or "."
    os.makedirs(result_dir, exist_ok=True)
    charts_root = os.path.join(result_dir, "charts")
    os.makedirs(charts_root, exist_ok=True)

    thresholds = {"tau_abs": args.tau_abs, "tau_rel": args.tau_rel, "tau_delta": args.tau_delta,
                  "tau_corr": args.tau_corr, "tau_cspi": args.tau_cspi}

    run_info = {}
    if ext in (".xlsx", ".xls", ".xlsm"):
        run_info, all_results = run_xlsx_pipeline(input_path, args, thresholds, charts_root)
    elif ext == ".csv":
        all_results = run_csv_pipeline(input_path, args, thresholds, charts_root)
    else:
        print(f"ERROR: unsupported input file type '{ext}'; expected .xlsx/.xls or .csv")
        sys.exit(1)

    if not all_results:
        print("No concept/stage produced usable results - nothing to write.")
        sys.exit(1)

    write_workbook(output_xlsx, all_results, run_info, args)
    print(f"\nDone. Bottleneck report written to: {output_xlsx}")
    print(f"Charts saved under: {charts_root}")


if __name__ == "__main__":
    main()
